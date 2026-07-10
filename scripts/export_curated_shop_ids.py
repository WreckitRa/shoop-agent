#!/usr/bin/env python3
"""
Export Shopify GIDs from pre-pass CHECKED workbook(s) into
src/lib/shopify/curated-shop-ids.ts for UCP filters.shop_ids.

Unions GIDs across every sheet that has a "Shopify GID" column, and across
all workbooks passed on the CLI (so new target lists can be *added* without
dropping the first-250 allowlist).

Usage:
  source .venv-prepass/bin/activate
  python scripts/export_curated_shop_ids.py Shoop-First-250-Shops-CHECKED.xlsx
  python scripts/export_curated_shop_ids.py \\
    Shoop-First-250-Shops-CHECKED.xlsx \\
    Shoop-Shoes-Accessories-Target-Lists-CHECKED.xlsx
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "src" / "lib" / "shopify" / "curated-shop-ids.ts"
GID_RE = re.compile(r"^gid://shopify/Shop/\d+$")


def collect_gids(workbook: Path) -> set[str]:
    wb = load_workbook(workbook, read_only=True, data_only=True)
    gids: set[str] = set()
    for sn in wb.sheetnames:
        ws = wb[sn]
        rows = ws.iter_rows(min_row=1, max_row=1, values_only=True)
        header_row = next(rows, None)
        if not header_row:
            continue
        headers = list(header_row)
        if "Shopify GID" not in headers:
            continue
        gi = headers.index("Shopify GID")
        for row in ws.iter_rows(min_row=2, values_only=True):
            if gi >= len(row):
                continue
            gid = row[gi]
            if gid and GID_RE.match(str(gid).strip()):
                gids.add(str(gid).strip())
    wb.close()
    return gids


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument(
        "workbooks",
        nargs="*",
        default=[str(ROOT / "Shoop-First-250-Shops-CHECKED.xlsx")],
        help="One or more CHECKED workbooks (GIDs are unioned)",
    )
    a = ap.parse_args()

    gids: set[str] = set()
    sources: list[str] = []
    for raw in a.workbooks:
        path = Path(raw)
        if not path.is_file():
            print(f"missing workbook: {path}", file=sys.stderr)
            sys.exit(1)
        found = collect_gids(path)
        print(f"{path.name}: {len(found)} gids")
        gids |= found
        sources.append(path.name)

    sorted_gids = sorted(gids, key=lambda g: (len(g.split("/")[-1]), g.split("/")[-1]))
    source_note = " + ".join(sources)
    lines = [
        "/**",
        " * Curated Shopify shop allowlist for Global Catalog `filters.shop_ids`.",
        " *",
        f" * Generated from {source_note} (pre-pass GIDs).",
        f" * {len(sorted_gids)} unique shops. Do not hand-edit — regenerate via:",
        " *   python scripts/export_curated_shop_ids.py",
        " */",
        "",
        "export const CURATED_SHOP_IDS: readonly string[] = [",
    ]
    for g in sorted_gids:
        lines.append(f'  "{g}",')
    lines.extend(
        [
            "] as const;",
            "",
            "/** True when catalog search should restrict to {@link CURATED_SHOP_IDS}. */",
            "export function isCuratedShopAllowlistEnabled(): boolean {",
            '  const raw = process.env.CATALOG_SHOP_ALLOWLIST?.trim().toLowerCase();',
            '  if (raw === "0" || raw === "false" || raw === "off") return false;',
            "  return CURATED_SHOP_IDS.length > 0;",
            "}",
            "",
        ]
    )
    OUT.write_text("\n".join(lines))
    print(f"Wrote {OUT.relative_to(ROOT)} ({len(sorted_gids)} gids)")


if __name__ == "__main__":
    main()
